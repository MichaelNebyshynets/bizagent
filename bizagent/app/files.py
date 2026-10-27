"""Файлы, которые пользователь прикладывает к чату.

КАК ЭТО УСТРОЕНО
  • Файл хранится на диске: <upload_dir>/<company_id>/<user_id>/<file_id>/
        original.<ext>  исходный файл
        text.txt        извлечённый текст (его и читает агент)
        meta.json       имя, размер, категория, дата
  • file_id — случайный uuid4 (32 hex-символа). Имя файла от пользователя НИКОГДА не используется
    как путь на диске, только для показа. Поэтому имя вида «../../app/main.py» ничего не сломает.
  • Изоляция: путь строится из company_id и user_id того, кто делает запрос (из токена).
    Чужой файл выглядит как несуществующий (NotFound), как и чужой диалог.
  • Текст извлекается один раз при загрузке. Модель получает его только через инструмент
    read_uploaded_file кусками по несколько тысяч символов, а не целиком в промпт.

Поддерживаемые форматы: .txt .md .csv (UTF-8 или Windows-1251, как в выгрузках 1С), .xlsx, .docx, .pdf (только с текстовым слоем, не сканы).
"""
from __future__ import annotations

import io
import json
import re
import shutil
import uuid
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from .clock import Clock
from .context import Principal
from .errors import AgentError, NotFound

MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 МБ
MAX_TEXT_CHARS = 200_000  # больше агенту всё равно не прочитать за разумное число шагов
ALLOWED_EXTENSIONS = (".txt", ".md", ".csv", ".xlsx", ".docx", ".pdf")
_ID_RE = re.compile(r"[0-9a-f]{32}")


class FileTooLarge(AgentError):
    status = 413
    code = "file_too_large"


class UnsupportedFile(AgentError):
    status = 415
    code = "unsupported_file"


class UnreadableFile(AgentError):
    status = 422
    code = "unreadable_file"


@dataclass
class StoredFile:
    id: str
    company_id: int
    user_id: int
    filename: str
    ext: str
    size: int
    chars: int
    truncated: bool
    category: str
    uploaded_at: str

    def public(self) -> dict[str, Any]:
        """То, что можно показать пользователю и модели (без company_id/user_id)."""
        return {
            "file_id": self.id,
            "filename": self.filename,
            "size": self.size,
            "chars": self.chars,
            "truncated": self.truncated,
            "category": self.category,
            "uploaded_at": self.uploaded_at,
        }


def clean_filename(name: str | None) -> str:
    """Имя только для показа: без пути, без управляющих символов, не длиннее 120 символов."""
    base = re.split(r"[\\/]", name or "")[-1]
    base = re.sub(r"[\x00-\x1f\x7f]", "", base).strip().strip(".") or "file"
    if len(base) > 120:
        stem, dot, ext = base.rpartition(".")
        base = (stem[: 110] + "." + ext) if dot and len(ext) <= 8 else base[:120]
    return base


# ------------------------------------------------------------------ извлечение текста


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):  # 1С и Excel в РБ часто сохраняют CSV в Windows-1251
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace("\n", " ").strip()


def _xlsx(data: bytes) -> str:
    try:
        import openpyxl
    except ImportError as e:
        raise UnsupportedFile("Для .xlsx на сервере нужен пакет openpyxl (pip install -r requirements.txt).") from e
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    size = 0
    try:
        for ws in wb.worksheets:
            out.append(f"## Лист: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [_cell(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                line = "; ".join(cells)
                out.append(line)
                size += len(line) + 1
                if size > MAX_TEXT_CHARS:
                    return "\n".join(out)
    finally:
        wb.close()
    return "\n".join(out)


def _docx(data: bytes) -> str:
    try:
        import docx  # пакет python-docx
    except ImportError as e:
        raise UnsupportedFile("Для .docx на сервере нужен пакет python-docx (pip install -r requirements.txt).") from e
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append("; ".join(_cell(c.text) for c in row.cells))
    return "\n".join(parts)


def _pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise UnsupportedFile("Для .pdf на сервере нужен пакет pypdf (pip install -r requirements.txt).") from e
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"## Страница {i}\n{text}")
        if sum(len(p) for p in pages) > MAX_TEXT_CHARS:
            break
    if not pages:
        raise UnreadableFile("В PDF нет текстового слоя (похоже на скан). Распознавание сканов пока не поддерживается.")
    return "\n\n".join(pages)


def extract_text(ext: str, data: bytes) -> str:
    try:
        if ext in (".txt", ".md", ".csv"):
            return _decode(data)
        if ext == ".xlsx":
            return _xlsx(data)
        if ext == ".docx":
            return _docx(data)
        if ext == ".pdf":
            return _pdf(data)
    except AgentError:
        raise
    except Exception as e:  # noqa: BLE001 — битый файл не должен ронять сервер
        raise UnreadableFile(f"Не удалось прочитать файл: {type(e).__name__}.") from e
    raise UnsupportedFile(f"Формат {ext} не поддерживается.")


# ------------------------------------------------------------------ хранилище


class FileStore:
    def __init__(self, root: str | Path, clock: Clock) -> None:
        self.root = Path(root)
        self._clock = clock

    def _dir(self, p: Principal, file_id: str) -> Path:
        return self.root / str(int(p.company_id)) / str(int(p.user_id)) / file_id

    def save(self, p: Principal, filename: str | None, data: bytes, category: str = "Все") -> StoredFile:
        name = clean_filename(filename)
        ext = Path(name).suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            raise UnsupportedFile(f"Формат «{ext or 'без расширения'}» не поддерживается. Можно: {', '.join(ALLOWED_EXTENSIONS)}.")
        if len(data) > MAX_FILE_BYTES:
            raise FileTooLarge(f"Файл больше {MAX_FILE_BYTES // (1024 * 1024)} МБ.")
        if not data:
            raise UnreadableFile("Файл пустой.")

        text = extract_text(ext, data).replace("\x00", "").strip()
        if not text:
            raise UnreadableFile("В файле не найден текст.")
        truncated = len(text) > MAX_TEXT_CHARS
        text = text[:MAX_TEXT_CHARS]

        sf = StoredFile(
            id=uuid.uuid4().hex, company_id=p.company_id, user_id=p.user_id, filename=name, ext=ext,
            size=len(data), chars=len(text), truncated=truncated, category=(category or "Все")[:40],
            uploaded_at=self._clock().isoformat(timespec="seconds"),
        )
        d = self._dir(p, sf.id)
        d.mkdir(parents=True, exist_ok=False)
        (d / f"original{ext}").write_bytes(data)
        (d / "text.txt").write_text(text, encoding="utf-8")
        (d / "meta.json").write_text(json.dumps(asdict(sf), ensure_ascii=False), encoding="utf-8")
        return sf

    def get(self, p: Principal, file_id: str) -> StoredFile:
        # Проверка формата id до обращения к диску: «../..» и прочее сразу NotFound
        if not isinstance(file_id, str) or not _ID_RE.fullmatch(file_id):
            raise NotFound("Файл не найден.")
        meta = self._dir(p, file_id) / "meta.json"
        if not meta.is_file():
            raise NotFound("Файл не найден.")
        sf = StoredFile(**json.loads(meta.read_text(encoding="utf-8")))
        if sf.company_id != p.company_id or sf.user_id != p.user_id:  # вторая линия защиты
            raise NotFound("Файл не найден.")
        return sf

    def read_text(self, p: Principal, file_id: str, offset: int, limit: int) -> tuple[StoredFile, str]:
        sf = self.get(p, file_id)
        text = (self._dir(p, file_id) / "text.txt").read_text(encoding="utf-8")
        return sf, text[offset: offset + limit]

    def list(self, p: Principal) -> list[StoredFile]:
        base = self.root / str(int(p.company_id)) / str(int(p.user_id))
        if not base.is_dir():
            return []
        out = []
        for d in base.iterdir():
            if d.is_dir() and _ID_RE.fullmatch(d.name):
                try:
                    out.append(self.get(p, d.name))
                except (NotFound, ValueError, TypeError):
                    continue
        return sorted(out, key=lambda f: f.uploaded_at, reverse=True)

    def delete(self, p: Principal, file_id: str) -> None:
        self.get(p, file_id)  # NotFound, если чужой или нет
        shutil.rmtree(self._dir(p, file_id))
